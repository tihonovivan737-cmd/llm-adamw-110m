"""Export training metrics.jsonl to a formatted Excel workbook."""
import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


TRAIN_COLUMNS = (
    ('step', 'Шаг'),
    ('tokens_seen', 'Токенов обработано'),
    ('step_tokens', 'Токенов за шаг'),
    ('train_loss', 'Train loss'),
    ('lr', 'Learning rate'),
    ('grad_norm_before_clip', 'Норма градиента'),
    ('step_seconds', 'Время шага, с'),
    ('tokens_per_second', 'Токенов/с'),
    ('peak_vram_gb', 'Пик VRAM, ГБ'),
)
MEMORY_COLUMNS = (
    ('step', 'Шаг'),
    ('gpu', 'GPU'),
    ('device_used_gb', 'Используется на GPU, ГБ'),
    ('device_total_gb', 'Всего на GPU, ГБ'),
    ('process_allocated_gb', 'Выделено процессом, ГБ'),
    ('process_reserved_gb', 'Зарезервировано процессом, ГБ'),
    ('step_peak_allocated_gb', 'Пик выделения за шаг, ГБ'),
    ('model_parameters_gb', 'Веса модели, ГБ'),
    ('gradients_gb', 'Градиенты, ГБ'),
    ('adamw_states_gb', 'Состояния AdamW, ГБ'),
    ('model_buffers_gb', 'Буферы модели, ГБ'),
    ('other_current_allocated_gb', 'Прочая текущая память, ГБ'),
    ('other_at_step_peak_approx_gb', 'Прочая память на пике, оценка, ГБ'),
)


def style_sheet(sheet, title, columns, table_name):
    sheet.title = title
    for index, (_, label) in enumerate(columns, start=1):
        sheet.cell(row=1, column=index, value=label)
    sheet.freeze_panes = 'A2'
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    sheet.row_dimensions[1].height = 34

    widths = {
        'Шаг': 10, 'Токенов обработано': 22, 'Токенов за шаг': 18,
        'Train loss': 14, 'Learning rate': 16, 'Норма градиента': 19,
        'Время шага, с': 16, 'Токенов/с': 16, 'Пик VRAM, ГБ': 16,
        'GPU': 9, 'Используется на GPU, ГБ': 23, 'Всего на GPU, ГБ': 18,
        'Выделено процессом, ГБ': 24, 'Зарезервировано процессом, ГБ': 29,
        'Пик выделения за шаг, ГБ': 25, 'Веса модели, ГБ': 18,
        'Градиенты, ГБ': 17, 'Состояния AdamW, ГБ': 21, 'Буферы модели, ГБ': 20,
        'Прочая текущая память, ГБ': 26,
        'Прочая память на пике, оценка, ГБ': 34,
    }
    for index, (_, label) in enumerate(columns, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = widths.get(label, 18)

    if sheet.max_row > 1:
        table = Table(displayName=table_name, ref=f'A1:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showFirstColumn=False,
                                               showLastColumn=False, showRowStripes=True,
                                               showColumnStripes=False)
        sheet.add_table(table)
        sheet.auto_filter.ref = table.ref
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.outlinePr.summaryBelow = True


def load_records(log_path):
    with log_path.open('r', encoding='utf-8') as source:
        for line_number, line in enumerate(source, start=1):
            if not line.strip():
                continue
            try:
                record = json.loads(line)
            except json.JSONDecodeError as error:
                raise ValueError(f'Invalid JSON on line {line_number}: {error}') from error
            if record.get('event') == 'train':
                yield record


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('logfile', type=Path, help='Run metrics.jsonl')
    parser.add_argument('--output', type=Path,
                        help='Output .xlsx path (defaults to training_metrics.xlsx beside the log)')
    args = parser.parse_args()
    if not args.logfile.is_file():
        parser.error(f'log file not found: {args.logfile}')

    output = args.output or args.logfile.with_name('training_metrics.xlsx')
    if output.suffix.lower() != '.xlsx':
        parser.error('output filename must end with .xlsx')
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    training = workbook.active
    style_sheet(training, 'Обучение', TRAIN_COLUMNS, 'TrainingSteps')
    memory = workbook.create_sheet('Память GPU')
    style_sheet(memory, 'Память GPU', MEMORY_COLUMNS, 'GpuMemory')

    train_number_formats = {
        'Train loss': '0.0000', 'Learning rate': '0.00000000',
        'Норма градиента': '0.000', 'Время шага, с': '0.00',
        'Токенов/с': '#,##0', 'Пик VRAM, ГБ': '0.000',
        'Токенов обработано': '#,##0', 'Токенов за шаг': '#,##0',
    }
    memory_number_formats = {label: '0.000' for _, label in MEMORY_COLUMNS if label != 'GPU'}
    for record in load_records(args.logfile):
        training.append([record.get(key) for key, _ in TRAIN_COLUMNS])
        for gpu_stats in record.get('gpu_memory', []):
            memory.append([record.get('step') if key == 'step' else gpu_stats.get(key)
                           for key, _ in MEMORY_COLUMNS])

    style_sheet(training, 'Обучение', TRAIN_COLUMNS, 'TrainingSteps')
    style_sheet(memory, 'Память GPU', MEMORY_COLUMNS, 'GpuMemory')
    for sheet, formats in ((training, train_number_formats), (memory, memory_number_formats)):
        for column_index, cell in enumerate(sheet[1], start=1):
            number_format = formats.get(cell.value)
            if number_format:
                for row in sheet.iter_rows(min_row=2, min_col=column_index, max_col=column_index):
                    row[0].number_format = number_format
    if training.max_row > 1:
        loss_col = next(index for index, (_, label) in enumerate(TRAIN_COLUMNS, start=1)
                        if label == 'Train loss')
        training.conditional_formatting.add(
            f'{training.cell(2, loss_col).coordinate}:{training.cell(training.max_row, loss_col).coordinate}',
            ColorScaleRule(start_type='min', start_color='C6EFD6',
                           mid_type='percentile', mid_value=50, mid_color='FFF2CC',
                           end_type='max', end_color='F4CCCC'))
    workbook.properties.title = 'Результаты обучения модели'
    workbook.properties.subject = 'Метрики шагов и память GPU'
    workbook.save(output)
    print(f'Excel-файл сохранён: {output} ({training.max_row - 1} шагов)')


if __name__ == '__main__':
    main()
